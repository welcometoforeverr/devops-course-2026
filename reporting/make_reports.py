"""Build practice reports from real repository pages and workflow evidence."""
from pathlib import Path
import html
import os
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from playwright.sync_api import sync_playwright, TimeoutError as PlaywrightTimeout

BASE = 'https://github.com/welcometoforeverr/devops-course-2026'
RUN3 = '40258811782'
RUN4 = os.environ.get('P4_RUN', '40259107686')
OUT = Path('report-output')
SC = Path('report-screens')
EV = Path('evidence4')
OUT.mkdir(exist_ok=True)
SC.mkdir(exist_ok=True)

def capture_sources():
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True)
        page = browser.new_page(viewport={'width': 1280, 'height': 820}, device_scale_factor=1)
        for name, url in {
            'dev': BASE+'/tree/dev',
            'workflow': BASE+'/blob/dev/.github/workflows/python-ci.yml',
            'pr3': BASE+'/pull/5',
            'ci3': BASE+'/actions/runs/'+RUN3,
            'ci4': BASE+'/actions/runs/'+RUN4,
            'proto': BASE+'/blob/feature/grpc-mtls/practice4/proto/api.proto',
            'nginx': BASE+'/blob/feature/grpc-mtls/practice4/nginx.conf',
            'django': BASE+'/blob/feature/grpc-mtls/practice4/client/urls.py',
        }.items():
            page.goto(url, wait_until='domcontentloaded', timeout=60000)
            if '/blob/' in url:
                page.get_by_role('textbox',name='file content').wait_for(state='visible',timeout=30000)
            try:
                page.wait_for_load_state('networkidle',timeout=5000)
            except PlaywrightTimeout:
                pass
            page.screenshot(path=str(SC/(name+'.png')))
        for name, source in [('certificates','certificates.txt'),('build','build.txt'),
                              ('checks','integration.txt')]:
            text=(EV/source).read_text()
            if name == 'certificates':
                text='\n'.join(text.splitlines()[-16:])
            if name == 'build':
                text='\n'.join(text.splitlines()[-18:])
            if name == 'checks':
                selected=[]
                for line in text.splitlines():
                    match=re.search(r'(DIRECT mTLS:|NGINX mTLS:|REJECTED without|CONCURRENT:|DJANGO:|HTTP 800|Ran 6 tests|^OK$)',line)
                    if match:selected.append(line[match.start():])
                text='\n'.join(selected)
            content='<html lang="ru"><meta charset="utf-8"><style>body{font:19px Arial;margin:40px;color:#172333}h1{font-size:25px}pre{font:16px/1.55 monospace;white-space:pre-wrap;overflow-wrap:anywhere;border:1px solid #ccd3da;padding:24px}</style>'
            content+='<h1>Журнал реального запуска GitHub Actions</h1><p>Ubuntu 24.04 · run '+RUN4+' · '+source+'</p><pre>'+html.escape(text)+'</pre>'
            page.set_content(content)
            page.screenshot(path=str(SC/(name+'.png')),full_page=True)
        browser.close()

def para(doc,text='',bold=False,align=None):
    p=doc.add_paragraph()
    if align is not None:p.alignment=align
    run=p.add_run(text);run.bold=bold
    return p

def code(doc,text):
    p=doc.add_paragraph()
    p.paragraph_format.first_line_indent=Cm(0)
    p.paragraph_format.space_after=Pt(8)
    p.paragraph_format.line_spacing=1
    for i,line in enumerate(text.strip().splitlines()):
        if i:p.add_run('\n')
        r=p.add_run(line);r.font.name='Liberation Mono';r.font.size=Pt(9)
    return p

def image(doc,path,caption,width=16.0):
    from PIL import Image
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after=Pt(3)
    with Image.open(path) as im:
        w,h=im.size
    # Keep each figure short enough to remain together with its caption.
    scale=min(width/w,10.0/h)
    p.add_run().add_picture(str(path),width=Cm(w*scale))
    p.paragraph_format.keep_with_next=True
    p=para(doc,caption,align=WD_ALIGN_PARAGRAPH.CENTER)
    p.paragraph_format.first_line_indent=Cm(0)
    for r in p.runs:r.font.size=Pt(10);r.italic=True

def table(doc,headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    for c,w in zip(t.columns,widths):c.width=Cm(w)
    for c,s in zip(t.rows[0].cells,headers):c.text=s
    for row in rows:
        for c,s in zip(t.add_row().cells,row):c.text=str(s)
    for i,row in enumerate(t.rows):
        for j,c in enumerate(row.cells):
            c.width=Cm(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcPr=c._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for edge in ('top','left','bottom','right'):
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            tcPr.append(borders)
            margins=OxmlElement('w:tcMar')
            for edge in ('top','left','bottom','right'):
                el=OxmlElement('w:'+edge);el.set(qn('w:w'),'90');el.set(qn('w:type'),'dxa');margins.append(el)
            tcPr.append(margins)
            if i==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E7EFF6');tcPr.append(shade)
            for p in c.paragraphs:
                p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.space_after=Pt(0)
                p.paragraph_format.line_spacing=1.08
                for r in p.runs:r.font.size=Pt(10.5);r.bold=(i==0)
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)

def heading(doc,text,new=False):
    if new:doc.add_page_break()
    doc.add_heading(text,1)

def init(num,topic):
    d=Document();s=d.sections[0]
    s.page_width=Cm(21);s.page_height=Cm(29.7)
    s.top_margin=Cm(2);s.bottom_margin=Cm(2);s.left_margin=Cm(2.5);s.right_margin=Cm(2)
    s.header_distance=Cm(0.8);s.footer_distance=Cm(0.8);s.different_first_page_header_footer=True
    for name in ('Normal','Title','Subtitle','Heading 1','Heading 2','Caption'):
        st=d.styles[name];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
    st=d.styles['Normal'];st.font.size=Pt(12)
    st.paragraph_format.line_spacing=1.15;st.paragraph_format.space_after=Pt(7)
    st.paragraph_format.widow_control=True
    for name in ('Heading 1','Heading 2'):
        st=d.styles[name];st.font.size=Pt(15 if name=='Heading 1' else 13)
        st.font.bold=True;st.paragraph_format.space_after=Pt(10)
    d.styles['Title'].font.size=Pt(19)
    d.core_properties.author='Калугин Никита Александрович'
    d.core_properties.title=f'Практическая работа {num} {topic}'
    for text in ('МИНОБРНАУКИ РОССИИ','МИРЭА — Российский технологический университет',
                 'Институт перспективных технологий', 'и индустриального программирования',
                 'Кафедра индустриального программирования'):
        para(d,text,align=WD_ALIGN_PARAGRAPH.CENTER)
    para(d,'\n\n')
    p=d.add_paragraph(f'Практическая работа {num}',style='Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    para(d,'по дисциплине «Инструменты девопс»',align=WD_ALIGN_PARAGRAPH.CENTER)
    para(d,topic,bold=True,align=WD_ALIGN_PARAGRAPH.CENTER)
    para(d,'\n\n')
    for text in ('Выполнил студент группы ЭФБО-12-24', 'Калугин Никита Александрович',
                 'Проверил', 'Гиматдинов Дамир Маратович'):
        para(d,text,align=WD_ALIGN_PARAGRAPH.RIGHT)
    para(d,'\n\nМосква 2026',align=WD_ALIGN_PARAGRAPH.CENTER)
    foot=s.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
    return d

def report3():
    d=init(3,'Ветвление Git и автоматическая проверка проекта')
    heading(d,'1 Цель работы и ветка разработки',True)
    para(d,'Цель работы — настроить разработку через ветку dev и автоматически проверять изменения при создании pull request. Для существующего Python-проекта настроен GitHub Actions; проверочный pull request успешно прошёл тесты и статический анализ.')
    para(d,'Репозиторий проекта: '+BASE)
    para(d,'В исходном файле задания работа обозначена как №2.2. В отчёте используется номер 3 по имени предоставленной методички.')
    para(d,'В локальном клоне создана ветка dev от main. Удалённая ветка dev опубликована в GitHub. Конфигурация проверок и тесты добавлены в dev до создания тематической ветки.')
    code(d,'git clone '+BASE+'.git\ncd devops-course-2026\ngit switch -c dev\n# Публикация ветки dev на GitHub')
    image(d,SC/'dev.png','Рисунок 1 — опубликованная ветка dev и файлы проекта')
    para(d,'Основной исходный модуль — calculator.py. В нём сохранены операции сложения и вычитания; добавлены строки документации для проверки Pylint. В tests размещены автоматические проверки арифметики.')
    heading(d,'2 Настройка GitHub Actions',True)
    para(d,'В .github/workflows/python-ci.yml создан workflow Python CI. Событие pull_request фильтруется по целевой ветке dev. Дополнительно проверки выполняются после push в dev. Задание запускается на ubuntu-24.04 с Python 3.12.')
    code(d,'on:\n  pull_request:\n    branches: [dev]\n  push:\n    branches: [dev]')
    image(d,SC/'workflow.png','Рисунок 2 — конфигурация workflow в ветке dev')
    para(d,'Порядок выполнения: получение исходников, настройка Python, установка Pylint 3.3.8, запуск unittest, анализ кода и загрузка lint.txt как артефакта pylint-report.')
    code(d,'python -m unittest discover -s tests -v\nset -o pipefail\npython -m pylint calculator.py tests --fail-under=9.0 | tee lint.txt')
    para(d,'В отличие от примера с «|| true», ошибка анализатора не подавляется: неуспешная проверка делает pipeline неуспешным. Отчёт загружается с if: always(), в том числе при ошибке проверки.')
    heading(d,'3 Тематическая ветка и pull request',True)
    para(d,'От подготовленной dev создана ветка feature/calculator-multiply. В calculator.py добавлена функция multiply. Тесты проверяют положительные и отрицательные множители, а также умножение на ноль.')
    code(d,'def multiply(a, b):\n    """Return the product of two numbers."""\n    return a * b')
    para(d,'Изменения опубликованы отдельным коммитом e7e16e7cf62f632be5adc20501d29e1e61c162c2. Создан pull request №5 из feature/calculator-multiply в dev.')
    para(d,'Pull request: '+BASE+'/pull/5')
    image(d,SC/'pr3.png','Рисунок 3 — pull request с целевой веткой dev')
    para(d,'Pull request содержит изменения трёх файлов: calculator.py, tests/test_calculator.py и practice3/README.md. Его открытие автоматически вызвало Python CI по событию pull_request.')
    heading(d,'4 Результаты проверки и вывод',True)
    image(d,SC/'ci3.png','Рисунок 4 — успешный запуск Python CI для pull request №5')
    table(d,['Проверка','Фактический результат'],[
        ('Событие запуска','pull_request, целевая ветка dev'),
        ('unittest','3 теста, OK'),('Pylint','10.00/10, порог 9.0'),
        ('GitHub Actions','Success, запуск 40258811782'),
        ('Артефакт','pylint-report, файл lint.txt')],[5.1,11.4])
    para(d,'Запуск с результатами: '+BASE+'/actions/runs/'+RUN3)
    para(d,'Цель достигнута: изменения из тематической ветки автоматически проверяются перед включением в dev. Pull request оставлен открытым с успешной проверкой, чтобы сохранить доступ к результату. Слияние в main для этой работы не требуется.')
    para(d,'Источник по настройке событий workflow: https://docs.github.com/en/actions/reference/workflows-and-actions/events-that-trigger-workflows')
    d.save(OUT/'Практическая работа 3.docx')

def report4():
    d=init(4,'Взаимодействие Python и Go через gRPC с mTLS и nginx')
    heading(d,'1 Цель работы и подготовка окружения',True)
    para(d,'Цель работы — организовать вызов Go-сервиса из Python/Django по gRPC, настроить взаимную TLS-аутентификацию и проксирование через nginx. Стенд выполнен и проверен на Ubuntu 24.04 в GitHub Actions.')
    para(d,'Исходный код: '+BASE+'/tree/feature/grpc-mtls/practice4')
    para(d,'Pull request: '+BASE+'/pull/6')
    table(d,['Компонент','Параметры стенда'],[
        ('Операционная система','Ubuntu 24.04, GitHub-hosted runner'),
        ('Python и Django','Python 3.12, Django 5.2.7'),
        ('gRPC Python','grpcio и grpcio-tools 1.71.0'),
        ('Go','Go 1.24.0; grpc 1.70.0; protobuf 1.36.5'),
        ('PKI и прокси','CFSSL и nginx из репозитория Ubuntu')],[5.0,11.5])
    code(d,'sudo apt-get update\nsudo apt-get install -y nginx golang-cfssl protobuf-compiler\ncd practice4\npython3 -m venv .venv\nsource .venv/bin/activate\npip install -r requirements.txt')
    para(d,'Структура проекта: certs — параметры PKI; proto — контракт сервиса; server — Go-сервис; client — Django и Python gRPC-клиент. generate.sh создаёт код из Protobuf, verify.sh запускает комплексную проверку.')
    para(d,'Прямой маршрут: Django → mTLS 127.0.0.1:8443 → Go. Маршрут через прокси: Django → mTLS 127.0.0.1:9448 → nginx → h2c 127.0.0.1:50051 → Go. Открытый h2c-порт доступен только через loopback.')
    heading(d,'2 Создание и проверка TLS сертификатов',True)
    para(d,'Создан собственный центр сертификации CA с ключом ECDSA P-256. Профиль server содержит назначение server auth, профиль client — client auth. Серверный сертификат включает SAN для 127.0.0.1 и localhost.')
    code(d,'bash certs.sh\nopenssl verify -CAfile certs/ca.pem -purpose sslserver certs/server.pem\nopenssl verify -CAfile certs/ca.pem -purpose sslclient certs/client.pem')
    image(d,SC/'certificates.png','Рисунок 1 — журнал CFSSL и успешная проверка обоих сертификатов')
    para(d,'Файлы ca.pem, server.pem и client.pem используются для построения доверия. Приватные ключи создаются с umask 077 и исключены из Git. Сертификаты конечных узлов действуют 8760 часов; в CA-конфигурации задан срок по умолчанию 168 часов.')
    para(d,'Go-сервер использует tls.RequireAndVerifyClientCert. Python передаёт корневой сертификат, сертификат клиента и его закрытый ключ в grpc.ssl_channel_credentials. Проверка серверного сертификата включена; подмена имени сервера не используется.')
    heading(d,'3 Контракт Protobuf и Go сервис',True)
    para(d,'В proto/api.proto определён DiceService с методом RollDie. Запрос не содержит полей, ответ RollDieResponse содержит целое поле value. Сервер возвращает число от 1 до 6 включительно.')
    image(d,SC/'proto.png','Рисунок 2 — общий контракт Python и Go в api.proto')
    code(d,'bash generate.sh\n# Python: client/api/api_pb2.py и api_pb2_grpc.py\n# Go: server/api/api.pb.go и api_grpc.pb.go\ngo build -o dice-server ./server\ngo test -race ./server/... -v')
    para(d,'В go_package указан путь python_go_grpc/server/api. Генератор Python исправляет абсолютный импорт api_pb2 на относительный. Версии protoc-gen-go и protoc-gen-go-grpc закреплены, чтобы повторный запуск давал совместимый код.')
    para(d,'Go-сервис открывает два слушателя: защищённый 8443 и локальный h2c 50051. gRPC-interceptor записывает метод, адрес клиента, статус и длительность вызова. Доступ к генератору случайных чисел защищён mutex; тест с 100 одновременными вызовами выполняется с детектором гонок.')
    heading(d,'4 Клиент Python и приложение Django',True)
    para(d,'Python-клиент создаёт secure_channel, вызывает DiceServiceStub.RollDie и закрывает канал после запроса. Для RPC задан таймаут 5 секунд. Пути сертификатов вычисляются от расположения модуля, поэтому не зависят от текущей директории вызова.')
    image(d,SC/'django.png','Рисунок 3 — отдельные маршруты Django для двух вариантов связи')
    para(d,'Корневой маршрут показывает два независимых броска кубика. POST-запрос защищён CSRF. /grpc обращается напрямую к Go, /nginx — через прокси; оба возвращают числовое значение и символ грани. При ошибке RPC возвращается HTTP 502 с контролируемым сообщением.')
    code(d,'python manage.py check\npython manage.py runserver 127.0.0.1:8000 --noreload')
    para(d,'HTML-шаблон хранится в client/templates/index.html. Установленные пакеты Django и файлы виртуального окружения не изменяются. Проверка manage.py check завершилась сообщением «System check identified no issues».')
    heading(d,'5 Настройка nginx и общего запуска',True)
    para(d,'nginx слушает порт 9448 с TLS и HTTP/2. Директива ssl_verify_client on требует клиентский сертификат от локального CA. grpc_pass направляет вызовы на loopback-порт Go-сервера 50051.')
    image(d,SC/'nginx.png','Рисунок 4 — конфигурация nginx с обязательной проверкой клиента')
    code(d,'mkdir -p logs\nnginx -p "$PWD/" -c nginx.conf -t\n./dice-server\n# В отдельных терминалах из practice4:\nnginx -p "$PWD/" -c nginx.conf\npython manage.py runserver 127.0.0.1:8000 --noreload')
    para(d,'В методичке у Python указан порт nginx 9448, а в конфигурации nginx — 9445. В реализации оба компонента используют 9448. nginx -t подтвердил корректность конфигурации. Общая автоматическая проверка запускается командой bash verify.sh.')
    heading(d,'6 Проверка в браузере и несколько клиентов',True)
    para(d,'В Chromium открыто Django-приложение на порту 8000. После нажатия «Бросить кубики» получены реальные значения от двух RPC-маршрутов. Значения могут различаться: это независимые случайные броски.')
    image(d,EV/'django-dice.png','Рисунок 5 — работа Django через прямой и проксируемый gRPC вызов')
    para(d,'Дополнительно запущены отдельные Django-процессы на портах 8001 и 8002. Для каждого проверены /grpc, /nginx и POST в браузере. Все процессы используют общий Go-сервис и nginx.')
    code(d,'python manage.py runserver 127.0.0.1:8001 --noreload\npython manage.py runserver 127.0.0.1:8002 --noreload')
    para(d,'Порт передаётся позиционным аргументом runserver. Упомянутый в методичке флаг -p для этой команды не используется. verify.sh хранит PID собственных процессов и завершает их после проверки.')
    heading(d,'7 Интеграционные тесты',True)
    image(d,SC/'checks.png','Рисунок 6 — фактический журнал интеграционных проверок',width=16.5)
    table(d,['Проверка','Результат'],[
        ('Прямой mTLS и nginx mTLS','По 20 вызовов, все значения в диапазоне 1–6'),
        ('Без клиентского сертификата','Оба защищённых входа отклоняют RPC'),
        ('Параллельные RPC','40 вызовов, 8 рабочих потоков, без ошибок'),
        ('Django и CSRF','GET, POST, /grpc и /nginx — HTTP 200'),
        ('Три web-клиента','Порты 8000, 8001 и 8002, оба маршрута работают'),
        ('Go race detector','TestConcurrentRolls — PASS')],[5.3,11.2])
    heading(d,'8 Итог выполнения',True)
    image(d,SC/'ci4.png','Рисунок 7 — успешный запуск стенда на Ubuntu 24.04')
    para(d,'Проверено взаимодействие Python и Go по общему Protobuf-контракту. Прямое соединение и вход nginx используют взаимную TLS-аутентификацию. Запрос без сертификата отклоняется, корректный клиент получает значение кубика. Проверены параллельные обращения и три Django-процесса.')
    para(d,'Запуск с результатами: '+BASE+'/actions/runs/'+RUN4)
    para(d,'Артефакт practice4-evidence содержит журналы генерации сертификатов, сборки, Go-теста, проверки nginx, интеграционных тестов и реальные скриншоты Chromium. Приватные ключи в артефакт не включены.')
    para(d,'Для повторения: клонировать репозиторий, перейти в feature/grpc-mtls, открыть practice4, установить зависимости и выполнить bash verify.sh. Для ручного просмотра после подготовки запустить Go, nginx и Django в отдельных терминалах по командам раздела 5.')
    para(d,'Использованные источники',bold=True)
    para(d,'1. Методические указания «Практическая работа №4».')
    para(d,'2. gRPC Authentication — https://grpc.io/docs/guides/auth/')
    para(d,'3. gRPC Python Quick start — https://grpc.io/docs/languages/python/quickstart/')
    para(d,'4. nginx gRPC module — https://nginx.org/en/docs/http/ngx_http_grpc_module.html')
    d.save(OUT/'Практическая работа 4.docx')

if __name__=='__main__':
    capture_sources()
    report3()
    report4()

